"""Generuje tabele Excel i CSV z plików JSON pakietu.

Użycie (z katalogu UiPath_AI_Warsztat, wymaga openpyxl):
  python narzedzia_pakietu/make_tables.py

Wynik:
  dane/Northstar_dane.xlsx    dane dla uczestnika (zgłoszenia, znane problemy, statusy, sprawy, kategorie, przykłady, pary)
  testy/Northstar_testy.xlsx  oczekiwane wyniki i testy techniczne (tylko do oceny, nie do wiedzy agenta)
  dane/csv/*.csv, testy/csv/*.csv  te same arkusze jako CSV (UTF-8 z BOM, separator średnik - otwiera się poprawnie w polskim Excelu)
JSON pozostaje źródłem. Po każdej zmianie JSON uruchom skrypt ponownie.
"""
import csv, json, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = PatternFill('solid', fgColor='142F47')
WHITE = Font(bold=True, color='FFFFFF')


def load(path):
    return json.load(open(path, encoding='utf8'))


def joined(v):
    return ', '.join(v) if isinstance(v, list) else v


DANE = [
    ('Zgloszenia', 'dane/zgloszenia.json',
     [('requestId', 'requestId'), ('requesterEmail', 'requesterEmail'), ('subject', 'subject'), ('body', 'body'), ('service', 'service')],
     {'body': 70, 'subject': 26}),
    ('Przyklady_cwiczeniowe', 'dane/przyklady_cwiczeniowe.json',
     [('requestId', 'requestId'), ('requesterEmail', 'requesterEmail'), ('subject', 'subject'), ('body', 'body'), ('service', 'service'), ('purpose', 'cel przykładu')],
     {'body': 60, 'cel przykładu': 50}),
    ('Znane_problemy', 'dane/known_issues.json',
     [('issueId', 'issueId'), ('service', 'service'), ('description', 'description'), ('workaround', 'workaround'), ('status', 'status')],
     {'description': 60, 'workaround': 50}),
    ('Statusy_uslug', 'dane/service_status.json',
     [('service', 'service'), ('status', 'status'), ('message', 'message'), ('eta', 'eta'), ('asOf', 'asOf')],
     {'message': 50}),
    ('Sprawy', 'dane/tickets.json',
     [('ticketId', 'ticketId'), ('requesterEmail', 'requesterEmail'), ('status', 'status'), ('summary', 'summary'), ('lastUpdate', 'lastUpdate'), ('eta', 'eta')],
     {'summary': 36, 'lastUpdate': 50}),
    ('Kategorie', 'dane/categories.json',
     [('name', 'name'), ('description', 'description')],
     {'description': 80}),
    ('Pary_podobienstwa', 'testy/pary_podobienstwa.json',
     [('id', 'id'), ('left', 'tekst A'), ('right', 'tekst B'), ('expected', 'oczekiwane'), ('note', 'uwagi'), (None, 'wynik score (wpisz)')],
     {'tekst A': 55, 'tekst B': 55, 'uwagi': 40}),
]

TESTY = [
    ('Oczekiwane_wyniki', 'testy/oczekiwane_wyniki.json',
     [('testId', 'testId'), ('requestId', 'requestId'), ('split', 'evaluation set'), ('expectedCategory', 'kategoria'), ('expectedNextAction', 'nextAction'),
      ('expectedHumanReview', 'humanReview'), ('expectedMatchedIssueId', 'matchedIssueId'), ('requiredTools', 'wymagane narzędzia'), ('note', 'uwagi'),
      (None, 'wynik v1 (wpisz)'), (None, 'wynik v2 (wpisz)')],
     {'uwagi': 50}),
    ('Testy_techniczne', 'testy/testy_techniczne.json',
     [('id', 'id'), ('inputId', 'zgłoszenie'), ('setup', 'przygotowanie'), ('expected', 'oczekiwany wynik'), ('mode', 'tryb'), (None, 'wynik (wpisz)')],
     {'przygotowanie': 55, 'oczekiwany wynik': 55, 'tryb': 36}),
    ('Rozmowy_testowe', 'testy/rozmowy_testowe.json',
     [('id', 'id'), ('requesterEmail', 'requesterEmail'), ('tura1', 'tura 1'), ('tura2', 'tura 2'), ('oczekiwane', 'oczekiwane zachowanie'),
      ('sprawdza', 'co sprawdza'), (None, 'wynik (wpisz)')],
     {'tura 1': 40, 'tura 2': 40, 'oczekiwane zachowanie': 60, 'co sprawdza': 30}),
]


def rows_of(path, cols):
    return [[joined(item.get(k, '')) if k else '' for k, _ in cols] for item in load(path)]


def build(sheets, xlsx, csvdir):
    wb = Workbook()
    wb.remove(wb.active)
    os.makedirs(csvdir, exist_ok=True)
    for name, path, cols, widths in sheets:
        headers = [h for _, h in cols]
        rows = rows_of(path, cols)
        ws = wb.create_sheet(name)
        ws.append(headers)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.fill, c.font = NAVY, WHITE
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(h, max(12, len(h) + 4))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        with open(os.path.join(csvdir, name + '.csv'), 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f, delimiter=';')
            w.writerow(headers)
            w.writerows(rows)
    wb.save(xlsx)
    print('ok', xlsx, [s[0] for s in sheets])


TOOLS = ['FindKnownIssue', 'GetServiceStatus', 'GetTicketStatus']


def behavior(tools):
    # Opis dla evaluatora Trajectory: tylko dobór narzędzi. Wynik końcowy sprawdzają evaluatory Exact match,
    # a historia przebiegu, którą widzi Trajectory, nie zawiera końcowej odpowiedzi agenta.
    other = [t for t in TOOLS if t not in tools]
    if tools:
        once = 'najwyżej raz' if len(tools) == 1 else 'każde najwyżej raz'
        return 'Agent wywołuje %s, %s. Nie wywołuje %s. Może sięgnąć do wiedzy Northstar_KB.' % (' i '.join(tools), once, ' ani '.join(other))
    return 'Agent nie wywołuje żadnego narzędzia procesowego (%s). Może sięgnąć do wiedzy Northstar_KB.' % ', '.join(TOOLS)


def eval_paste(split, out):
    """Bloki do wklejenia w Studio Web → Evaluations: Input, Expected output, Expected agent behavior."""
    tickets = {t['requestId']: t for t in load('dane/zgloszenia.json')}
    nl = chr(10)
    parts = ['PRZYPADKI DO EVALUATION SETU: Northstar %s' % split,
             'W Studio Web: agent SupportAgent → Evaluations → evaluation set "Northstar %s" → dla każdego zgłoszenia dodaj przypadek i wklej Input, Expected output i Expected agent behavior.' % split,
             'Plik jest generowany z dane/zgloszenia.json i testy/oczekiwane_wyniki.json (narzedzia_pakietu/make_tables.py).', '']
    for e in load('testy/oczekiwane_wyniki.json'):
        if e['split'] != split:
            continue
        t = tickets[e['requestId']]
        inp = {k: t[k] for k in ('requestId', 'requesterEmail', 'subject', 'body', 'service')}
        exp = {'requestId': e['requestId'], 'category': e['expectedCategory'], 'nextAction': e['expectedNextAction'],
               'humanReview': e['expectedHumanReview'], 'matchedIssueId': e.get('expectedMatchedIssueId') or ''}
        tools = e.get('requiredTools') or []
        beh = behavior(tools)
        parts += ['=' * 60, e['requestId'], '=' * 60, 'Input:', json.dumps(inp, ensure_ascii=False, indent=2), '',
                  'Expected output:', json.dumps(exp, ensure_ascii=False, indent=2), '', 'Expected agent behavior:', beh, '']
    open(out, 'w', encoding='utf8', newline='\n').write(nl.join(parts))
    print('ok', out)


def eval_import(split, out):
    """Plik do Studio Web → Evaluations → Import (format z szablonu Studio Web). UUID stałe, żeby plik nie zmieniał się przy regeneracji."""
    import uuid
    ns = uuid.UUID('6b2f6a3e-1d3c-4f59-9d0b-6e4f2a1c7b10')
    set_id = str(uuid.uuid5(ns, 'set-' + split))
    stamp = '2026-10-09T00:00:00.000Z'
    tickets = {t['requestId']: t for t in load('dane/zgloszenia.json')}
    evals = []
    for e in load('testy/oczekiwane_wyniki.json'):
        if e['split'] != split:
            continue
        t = tickets[e['requestId']]
        tools = e.get('requiredTools') or []
        beh = behavior(tools)
        evals.append({
            'id': str(uuid.uuid5(ns, split + '-' + e['requestId'])), 'name': e['requestId'],
            'inputs': {k: t[k] for k in ('requestId', 'requesterEmail', 'subject', 'body', 'service')},
            'expectedOutput': {'requestId': e['requestId'], 'category': e['expectedCategory'], 'nextAction': e['expectedNextAction'],
                               'humanReview': e['expectedHumanReview'], 'matchedIssueId': e.get('expectedMatchedIssueId') or ''},
            'expectedAgentBehavior': beh, 'simulationInstructions': '', 'simulateInput': False, 'inputGenerationInstructions': '',
            'simulateTools': False, 'toolsToSimulate': [], 'evalSetId': set_id, 'createdAt': stamp, 'updatedAt': stamp})
    doc = {'fileName': os.path.basename(out), 'id': set_id, 'name': 'Northstar ' + split, 'batchSize': 10, 'evaluatorRefs': [],
           'evaluations': evals, 'modelSettings': [], 'createdAt': stamp, 'updatedAt': stamp,
           'agentMemoryEnabled': False, 'agentMemorySettings': [], 'lineByLineEvaluation': False}
    json.dump(doc, open(out, 'w', encoding='utf8', newline='\n'), ensure_ascii=False, indent=2)
    print('ok', out, len(evals))


def eval_import_csv(split, out):
    """Wariant CSV importu (szablon Studio Web: name, <pola wejścia>, expected_output). Brak kolumny na Expected agent behavior."""
    tickets = {t['requestId']: t for t in load('dane/zgloszenia.json')}
    fields = ['requestId', 'requesterEmail', 'subject', 'body', 'service']
    with open(out, 'w', encoding='utf8', newline='') as f:
        w = csv.writer(f)
        w.writerow(['name'] + fields + ['expected_output'])
        for e in load('testy/oczekiwane_wyniki.json'):
            if e['split'] != split:
                continue
            t = tickets[e['requestId']]
            exp = {'requestId': e['requestId'], 'category': e['expectedCategory'], 'nextAction': e['expectedNextAction'],
                   'humanReview': e['expectedHumanReview'], 'matchedIssueId': e.get('expectedMatchedIssueId') or ''}
            w.writerow([e['requestId']] + [t[k] for k in fields] + [json.dumps(exp, ensure_ascii=False)])
    print('ok', out)


def known_issues_vb(out):
    # Wartość domyślna zmiennej lstKnownIssues w FindKnownIssue (ćw. 2): projekt chmurowy nie czyta plików z dysku
    q = lambda v: '"' + str(v).replace('"', '""') + '"'
    items = ['New Dictionary(Of String, String) From {' + ', '.join('{%s, %s}' % (q(k), q(r[k])) for k in ('issueId', 'service', 'description', 'workaround')) + '}'
             for r in load('dane/known_issues.json')]
    open(out, 'w', encoding='utf8', newline='\n').write('New List(Of Dictionary(Of String, String)) From {' + ', '.join(items) + '}\n')
    print('ok', out)


def tickets_for_bucket(out):
    # sprawy.csv do bucketu Northstar_Rejestr (ćw. 3A): bez BOM, średnik, LF; czyta je GetTicketStatus
    cols = ['ticketId', 'requesterEmail', 'status', 'summary', 'lastUpdate', 'eta']
    lines = [';'.join(cols)] + [';'.join(str(r[c]) for c in cols) for r in load('dane/tickets.json')]
    open(out, 'w', encoding='utf8', newline='\n').write('\n'.join(lines) + '\n')
    print('ok', out)


build(DANE, 'dane/Northstar_dane.xlsx', 'dane/csv')
known_issues_vb('dane/known_issues_do_wklejenia.txt')
tickets_for_bucket('dane/sprawy.csv')
build(TESTY, 'testy/Northstar_testy.xlsx', 'testy/csv')
eval_paste('DEVELOPMENT', 'testy/ewaluacja_do_wklejenia.txt')
eval_import('DEVELOPMENT', 'testy/import_Northstar_DEVELOPMENT.json')
eval_import('HOLDOUT', 'testy/import_Northstar_HOLDOUT.json')
eval_import_csv('DEVELOPMENT', 'testy/import_Northstar_DEVELOPMENT.csv')
