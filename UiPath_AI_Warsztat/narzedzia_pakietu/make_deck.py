"""Generuje prezentację UiPath_AI_Warsztaty.pptx z treści w deck_tresc.py.

Użycie (z katalogu UiPath_AI_Warsztat, wymaga python-pptx):
  python narzedzia_pakietu/make_deck.py
Treść slajdów i notatki prowadzącego są w narzedzia_pakietu/deck_tresc.py; tu tylko układ i styl.
Styl: DejaVu Sans, granat 142F47, turkus 087F8C, etykieta dnia, numer strony.
"""
import os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from deck_tresc import SLIDES  # noqa: E402

NAVY = RGBColor(0x14, 0x2F, 0x47)
TEAL = RGBColor(0x08, 0x7F, 0x8C)
TEAL_L = RGBColor(0x6A, 0xD1, 0xCE)
CREAM = RGBColor(0xF6, 0xF5, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x53, 0x68, 0x78)
MUTED_D = RGBColor(0xA9, 0xBC, 0xCB)
PALE = RGBColor(0xD8, 0xE4, 0xEB)
ZEBRA = RGBColor(0xED, 0xF1, 0xF3)
WARN = RGBColor(0xD9, 0x82, 0x2B)
WARN_L = RGBColor(0xFB, 0xEF, 0xE2)
FONT = 'DejaVu Sans'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333


def text(slide, x, y, w, h, s, size=18, bold=False, color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, line in enumerate(s.split('\n') if isinstance(s, str) else s):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.25):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.12
    return sh


def box_text(sh, title, sub=None, size=17, color=NAVY, subcolor=MUTED, subsize=13, align=PP_ALIGN.CENTER):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ('margin_left', 'margin_right'):
        setattr(tf, m, Inches(0.1))
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = title
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), True, color
    if sub:
        p = tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = sub
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(subsize), subcolor


def arrow(slide, x1, y1, x2, y2, color=TEAL):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)


def frame(s, dark=False):
    slide = prs.slides.add_slide(BLANK)
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = NAVY if dark else CREAM
    box(slide, 0, 0, 0.09, 7.5, fill=TEAL, line=None, shape=MSO_SHAPE.RECTANGLE)
    box(slide, 12.03, 0.4, 0.42, 0.06, fill=TEAL, line=None, shape=MSO_SHAPE.RECTANGLE)
    lab = s.get('label', '')
    if lab:
        text(slide, 0.8, 0.36, 8.3, 0.35, lab, 13.5, True, TEAL_L if dark else TEAL)
    text(slide, 11.75, 6.88, 0.8, 0.3, '%02d' % (len(prs.slides)), 12, False, MUTED_D if dark else MUTED, PP_ALIGN.RIGHT)
    return slide


def title(slide, s, dark=False, y=0.95, size=33):
    if len(s['title']) > 44:
        size = 28
    text(slide, 0.8, y, 11.8, 1.1, s['title'], size, True, WHITE if dark else NAVY)
    if s.get('subtitle'):
        text(slide, 0.85, y + 0.95, 11.6, 0.6, s['subtitle'], 17, False, TEAL_L if dark else TEAL)


def footer_note(slide, s, dark=False, y=6.3):
    if s.get('takeaway'):
        box(slide, 0.85, y + 0.02, 0.07, 0.42, fill=TEAL, line=None, shape=MSO_SHAPE.RECTANGLE)
        text(slide, 1.05, y, 11.3, 0.5, s['takeaway'], 17, True, TEAL_L if dark else NAVY, anchor=MSO_ANCHOR.MIDDLE)


def notes(slide, s):
    n = s.get('notes', '').strip()
    src = s.get('sources')
    if src:
        n += '\nŹródła: ' + ' ; '.join(src)
    slide.notes_slide.notes_text_frame.text = n


# --- typy slajdów -----------------------------------------------------------

def s_cover(s):
    slide = frame(s, dark=True)
    text(slide, 0.8, 2.1, 11.6, 1.6, s['title'], 52, True, WHITE)
    text(slide, 0.85, 3.85, 11.6, 0.7, s['subtitle'], 24, False, PALE)
    text(slide, 0.85, 5.45, 11.6, 0.5, s.get('author', ''), 19, False, TEAL_L)
    return slide


def s_section(s):
    slide = frame(s, dark=True)
    if s.get('kicker'):
        text(slide, 0.85, 1.75, 11.5, 0.5, s['kicker'], 17, True, TEAL_L)
    text(slide, 0.8, 2.25, 11.6, 1.5, s['title'], 44, True, WHITE)
    y = 4.15
    for line in s.get('lines', []):
        text(slide, 0.85, y, 11.3, 0.55, line, 21, False, PALE)
        y += 0.62
    return slide


def s_question(s):
    slide = frame(s, dark=True)
    text(slide, 0.85, 1.5, 11.5, 0.5, s.get('kicker', 'PYTANIE DO GRUPY'), 15, True, TEAL_L)
    text(slide, 0.8, 2.1, 11.6, 2.4, s['title'], 38, True, WHITE)
    y = 4.75
    for line in s.get('lines', []):
        text(slide, 0.85, y, 11.3, 0.5, line, 19, False, PALE)
        y += 0.55
    return slide


def s_bullets(s):
    dark = s.get('dark', False)
    slide = frame(s, dark)
    title(slide, s, dark)
    items = s['items']
    top = 2.55 if s.get('subtitle') else 2.3
    gap = min(1.05, (5.85 - top) / len(items))
    for i, it in enumerate(items):
        y = top + i * gap
        head, _, tail = it.partition(': ') if ': ' in it and s.get('heads', True) else ('', '', it)
        box(slide, 0.85, y + 0.08, 0.13, 0.13, fill=TEAL, line=None, shape=MSO_SHAPE.OVAL)
        tb = slide.shapes.add_textbox(Inches(1.15), Inches(y), Inches(11.2), Inches(gap))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        if head:
            r = p.add_run(); r.text = head + ': '
            r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(20), True, (WHITE if dark else NAVY)
        r = p.add_run(); r.text = tail
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(20), (PALE if dark else NAVY)
    footer_note(slide, s, dark)
    return slide


def s_table(s):
    slide = frame(s)
    title(slide, s)
    rows = s['rows']
    top = 2.55 if s.get('subtitle') else 2.35
    widths = s.get('widths') or [11.7 / len(rows[0])] * len(rows[0])
    rh = s.get('row_h', 0.62)
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.85), Inches(top), Inches(sum(widths)), Inches(rh * len(rows)))
    tbl = shape.table
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(rh)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (ZEBRA if i % 2 == 0 else WHITE)
            c.margin_left = c.margin_right = Inches(0.12)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            tf.paragraphs[0].text = ''
            r = tf.paragraphs[0].add_run(); r.text = val
            r.font.name = FONT
            r.font.size = Pt(s.get('font', 15) if i else 14)
            r.font.bold = (i == 0) or (j == 0 and s.get('bold_first', True))
            r.font.color.rgb = WHITE if i == 0 else (TEAL if (j == 0 and s.get('teal_first')) else NAVY)
    tblPr = shape._element.graphic.graphicData.tbl.tblPr
    tblPr.set('bandRow', '0'); tblPr.set('firstRow', '0')
    footer_note(slide, s, y=min(6.35, top + rh * len(rows) + 0.25))
    return slide


def s_flow(s):
    slide = frame(s)
    title(slide, s)
    rows = s['rows']  # lista wierszy; wiersz = (etykieta lub None, [(tytuł, podpis, styl), ...])
    top = 2.55 if s.get('subtitle') else 2.4
    rh = s.get('row_h', 1.15)
    gap_r = s.get('row_gap', 0.3)
    if len(rows) == 1:
        top = 3.05
    lw = s.get('label_w', 2.3)
    for ri, (lab, steps) in enumerate(rows):
        y = top + ri * (rh + gap_r)
        x0 = 0.85
        if lab:
            text(slide, 0.85, y, lw, rh, lab, s.get('label_size', 18), True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
            x0 = 0.85 + lw + 0.15
        n = len(steps)
        gapw = 0.55
        bw = (12.45 - x0 - gapw * (n - 1)) / n
        for k, st in enumerate(steps):
            t, sub, style = (list(st) + [None, None])[:3]
            x = x0 + k * (bw + gapw)
            fill, line, col, scol = WHITE, TEAL, NAVY, MUTED
            if style == 'ai':
                fill, line = RGBColor(0xE3, 0xF4, 0xF4), TEAL
            elif style == 'human':
                fill, line = WARN_L, WARN
            elif style == 'dark':
                fill, line, col, scol = NAVY, NAVY, WHITE, PALE
            elif style == 'muted':
                fill, line, col = ZEBRA, PALE, MUTED
            sh = box(slide, x, y, bw, rh, fill=fill, line=line, lw=1.5)
            box_text(sh, t, sub, size=s.get('font', 17), color=col, subcolor=scol)
            if k < n - 1:
                arrow(slide, x + bw + 0.05, y + rh / 2, x + bw + gapw - 0.05, y + rh / 2)
    if s.get('legend'):
        text(slide, 0.85, 5.75, 11.5, 0.4, s['legend'], 13, False, MUTED)
    footer_note(slide, s)
    return slide


def s_cards(s):
    dark = s.get('dark', False)
    slide = frame(s, dark)
    title(slide, s, dark)
    cards = s['cards']
    n = len(cards)
    cols = s.get('cols', n if n <= 4 else (3 if n in (5, 6) else 4))
    nrows = (n + cols - 1) // cols
    top = 2.55 if s.get('subtitle') else 2.35
    gw, gh = 0.3, 0.3
    avail = (6.05 if s.get('takeaway') else 6.6) - top
    ch = min(s.get('max_h', 2.5), (avail - gh * (nrows - 1)) / nrows)
    cw = (11.7 - gw * (cols - 1)) / cols
    for i, c in enumerate(cards):
        head, body = c[0], c[1]
        style = c[2] if len(c) > 2 else None
        r, k = divmod(i, cols)
        x, y = 0.85 + k * (cw + gw), top + r * (ch + gh)
        fill, line, hc, bc = WHITE, TEAL, TEAL, NAVY
        if dark:
            fill, line, hc, bc = RGBColor(0x1B, 0x3A, 0x55), RGBColor(0x2A, 0x50, 0x70), TEAL_L, WHITE
        if style == 'warn':
            fill, line, hc = WARN_L, WARN, WARN
        sh = box(slide, x, y, cw, ch, fill=fill, line=line)
        box(slide, x, y + 0.18, 0.07, min(0.6, ch - 0.36), fill=hc, line=None, shape=MSO_SHAPE.RECTANGLE)
        tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left, tf.margin_right, tf.margin_top = Inches(0.25), Inches(0.15), Inches(0.15)
        p = tf.paragraphs[0]
        rr = p.add_run(); rr.text = head
        rr.font.name, rr.font.size, rr.font.bold, rr.font.color.rgb = FONT, Pt(s.get('head_size', 19)), True, hc
        for line in body.split('\n'):
            p = tf.add_paragraph(); p.space_before = Pt(6)
            rr = p.add_run(); rr.text = line
            rr.font.name, rr.font.size, rr.font.color.rgb = FONT, Pt(s.get('font', 17)), bc
    footer_note(slide, s, dark)
    return slide


def s_chat(s):
    slide = frame(s)
    title(slide, s)
    y = 2.15
    for who, msg in s['turns']:
        user = who == 'Użytkownik'
        tool = who not in ('Użytkownik', 'Agent')
        w = 6.4
        x = 0.85 if user else (12.45 - w)
        fill = WHITE if user else (WARN_L if tool else RGBColor(0xE3, 0xF4, 0xF4))
        line = PALE if user else (WARN if tool else TEAL)
        sh = box(slide, x, y, w, 0.68, fill=fill, line=line)
        box_text(sh, who, msg, size=13, subsize=16, color=(WARN if tool else TEAL), subcolor=NAVY, align=PP_ALIGN.LEFT)
        y += 0.8
    footer_note(slide, s)
    return slide


def s_lab(s):
    slide = frame(s, dark=True)
    badge = box(slide, 0.85, 0.95, 1.05, 0.42, fill=TEAL, line=None)
    box_text(badge, 'LAB', size=13, color=WHITE)
    if s.get('time'):
        text(slide, 2.05, 0.98, 4, 0.4, s['time'], 15, True, TEAL_L, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 0.8, 1.55, 11.8, 1.0, s['title'], 33, True, WHITE)
    y = 2.75
    for lab, val in (('Budujesz', s['build']), ('Sprawdzasz', s['check'])):
        text(slide, 0.85, y, 2.2, 0.5, lab.upper(), 14, True, TEAL_L)
        text(slide, 3.0, y - 0.05, 9.4, 0.9, val, 21, False, WHITE)
        y += 1.05
    sh = box(slide, 0.85, 5.05, 11.6, 0.95, fill=RGBColor(0x1B, 0x3A, 0x55), line=TEAL)
    box_text(sh, 'Gotowe, gdy', s['done'], size=13, color=TEAL_L, subcolor=WHITE, subsize=19, align=PP_ALIGN.LEFT)
    text(slide, 0.85, 6.3, 11.5, 0.4, 'Kroki, dane i punkt kontrolny: instrukcja warsztatowa, ' + s['ref'], 14, False, MUTED_D)
    return slide


KINDS = {'cover': s_cover, 'section': s_section, 'question': s_question, 'bullets': s_bullets, 'table': s_table,
         'flow': s_flow, 'cards': s_cards, 'chat': s_chat, 'lab': s_lab}

for s in SLIDES:
    slide = KINDS[s['kind']](s)
    notes(slide, s)

out = os.path.join(os.path.dirname(HERE), 'UiPath_AI_Warsztaty.pptx')
prs.save(out)
print('ok', out, len(prs.slides))
